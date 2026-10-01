import calendar
from datetime import date
from pathlib import Path

from django.contrib import messages
from django.contrib.auth.mixins import LoginRequiredMixin
from django.core.exceptions import PermissionDenied, ValidationError
from django.db import transaction
from django.db.models import Count, Q
from django.http import FileResponse, HttpResponse, JsonResponse, QueryDict
from django.shortcuts import get_object_or_404, redirect
from django.urls import reverse_lazy
from django.utils import timezone
from django.views import View
from django.views.generic import CreateView, DetailView, ListView, TemplateView, UpdateView

from accounts.models import User
from organization.models import Department
from .forms import IssueNoteForm, MeetingSeriesForm, ProjectForm, TaskForm
from .excel_export import EXPORT_FIELDS, build_task_workbook
from .models import IssueNote, IssueNoteAttachment, Meeting, MeetingItem, MeetingSeries, Project, Tag, Task, TaskCategory
from .permissions import can_edit_task, department_scope_ids, visible_tasks


class TaskQuerysetMixin(LoginRequiredMixin):
    def get_queryset(self):
        queryset = Task.objects.select_related(
            "category", "request_department", "requestor", "owner_department", "primary_assignee", "created_by"
        ).prefetch_related("tags", "collaborators")
        return visible_tasks(self.request.user, queryset)


class MeetingCenterView(LoginRequiredMixin, TemplateView):
    template_name = "workitems/meeting_center.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        try:
            series = MeetingSeries.objects.filter(is_active=True).select_related(
                "project", "department", "chair", "recorder"
            ).prefetch_related("members", "meetings")
            cards = []
            for item in series:
                item.latest_meeting = item.meetings.order_by("-scheduled_at").first()
                item.open_tasks = self._series_tasks(item).exclude(status__in=[Task.Status.COMPLETED, Task.Status.CANCELLED]).count()
                cards.append(item)
            active_meetings = Meeting.objects.filter(status=Meeting.Status.IN_PROGRESS).select_related("series", "chair").order_by("-started_at")
            recent_meetings = Meeting.objects.filter(status=Meeting.Status.COMPLETED).select_related("series", "chair").order_by("-ended_at")[:8]
        except Exception:
            cards = []
            active_meetings = []
            recent_meetings = []
        context.update({
            "series_cards": cards,
            "active_meetings": active_meetings,
            "recent_meetings": recent_meetings,
            "can_manage": getattr(self.request.user, "role", None) != User.Role.VIEWER if self.request.user.is_authenticated else True,
        })
        return context

    def _series_tasks(self, series):
        tasks = visible_tasks(self.request.user, Task.objects.all())
        if series.project_id:
            return tasks.filter(project=series.project)
        if series.department_id:
            return tasks.filter(owner_department_id__in=department_scope_ids(series.department))
        return tasks.none()


class MeetingSeriesCreateView(LoginRequiredMixin, CreateView):
    model = MeetingSeries
    form_class = MeetingSeriesForm
    template_name = "workitems/meeting_series_form.html"
    success_url = reverse_lazy("meeting_center")

    def dispatch(self, request, *args, **kwargs):
        if request.user.is_authenticated and request.user.role == User.Role.VIEWER:
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def form_valid(self, form):
        form.instance.created_by = self.request.user
        messages.success(self.request, "會議系列已建立，可以直接開始會議。")
        return super().form_valid(form)


class MeetingStartView(LoginRequiredMixin, View):
    def post(self, request, pk):
        if request.user.role == User.Role.VIEWER:
            raise PermissionDenied
        series = get_object_or_404(MeetingSeries.objects.prefetch_related("members"), pk=pk, is_active=True)
        current = series.meetings.filter(status=Meeting.Status.IN_PROGRESS).first()
        if current:
            return redirect("meeting_detail", pk=current.pk)
        now = timezone.now()
        meeting = Meeting.objects.create(
            series=series, title=f"{series.name}｜{timezone.localdate():%Y/%m/%d}", scheduled_at=now,
            status=Meeting.Status.IN_PROGRESS, chair=series.chair, recorder=series.recorder,
            started_at=now, created_by=request.user,
        )
        meeting.attendees.set(series.members.all())
        tasks = visible_tasks(request.user, Task.objects.all()).exclude(status__in=[Task.Status.COMPLETED, Task.Status.CANCELLED])
        if series.project_id:
            tasks = tasks.filter(project=series.project)
        elif series.department_id:
            tasks = tasks.filter(owner_department_id__in=department_scope_ids(series.department))
        else:
            tasks = tasks.none()
        order = 0
        for task in tasks.select_related("primary_assignee").order_by("due_date", "task_no"):
            MeetingItem.objects.create(meeting=meeting, task=task, topic=task.title, owner=task.primary_assignee, due_date=task.due_date, order=order)
            order += 1
        if series.carry_unfinished:
            previous = series.meetings.filter(status=Meeting.Status.COMPLETED).exclude(pk=meeting.pk).order_by("-ended_at").first()
            if previous:
                existing_task_ids = set(meeting.items.exclude(task=None).values_list("task_id", flat=True))
                for old in previous.items.filter(carry_to_next=True).exclude(next_action=""):
                    if old.task_id and old.task_id in existing_task_ids:
                        continue
                    MeetingItem.objects.create(meeting=meeting, task=old.task, topic=f"延續追蹤｜{old.topic}", next_action=old.next_action, owner=old.owner, due_date=old.due_date, order=order)
                    order += 1
        messages.success(request, f"已開始「{meeting.title}」，並帶入 {meeting.items.count()} 個追蹤事項。")
        return redirect("meeting_detail", pk=meeting.pk)


class MeetingRecordView(LoginRequiredMixin, DetailView):
    model = Meeting
    template_name = "workitems/meeting_record.html"
    context_object_name = "meeting"

    def get_queryset(self):
        return Meeting.objects.select_related("series", "chair", "recorder").prefetch_related("attendees", "items__task", "items__owner")
        
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        from accounts.models import User
        context['active_users'] = User.objects.filter(is_active=True).order_by("display_name")
        return context

    def post(self, request, pk):
        if request.user.role == User.Role.VIEWER:
            raise PermissionDenied
        meeting = self.get_object()
        
        from datetime import datetime
        raw_date = request.POST.get("started_at", "")
        if raw_date:
            try:
                dt = datetime.strptime(raw_date, "%Y-%m-%d").date()
                if meeting.started_at:
                    meeting.started_at = meeting.started_at.replace(year=dt.year, month=dt.month, day=dt.day)
                else:
                    meeting.started_at = datetime(dt.year, dt.month, dt.day)
            except ValueError:
                pass
                
        meeting.location = request.POST.get("location", "").strip()
        meeting.chair_id = request.POST.get("chair") or meeting.chair_id
        meeting.recorder_id = request.POST.get("recorder") or meeting.recorder_id
        meeting.save()
        
        attendee_ids = request.POST.getlist("attendees")
        if attendee_ids:
            meeting.attendees.set(attendee_ids)
            
        messages.success(request, "會議紀錄屬性已更新")
        return redirect("meeting_record", pk=meeting.pk)

class MeetingDetailView(LoginRequiredMixin, DetailView):
    model = Meeting
    template_name = "workitems/meeting_detail.html"
    context_object_name = "meeting"

    def get_queryset(self):
        return Meeting.objects.select_related("series", "chair", "recorder").prefetch_related("attendees", "items__task", "items__owner")

    def post(self, request, pk):
        if request.user.role == User.Role.VIEWER:
            raise PermissionDenied
        meeting = self.get_object()
        action = request.POST.get("action")
        if action == "add_item":
            topic = request.POST.get("topic", "").strip()
            if topic:
                MeetingItem.objects.create(meeting=meeting, topic=topic, order=meeting.items.count())
                messages.success(request, "已加入臨時討論事項。")
        elif action == "save_summary":
            meeting.summary = request.POST.get("summary", "").strip()
            meeting.save(update_fields=["summary"])
            messages.success(request, "會議摘要已儲存。")
        elif action == "update_meeting_info":
            from datetime import datetime
            raw_date = request.POST.get("started_at", "")
            if raw_date:
                try:
                    dt = datetime.strptime(raw_date, "%Y-%m-%d").date()
                    if meeting.started_at:
                        meeting.started_at = meeting.started_at.replace(year=dt.year, month=dt.month, day=dt.day)
                    else:
                        meeting.started_at = datetime(dt.year, dt.month, dt.day)
                except ValueError:
                    pass
            meeting.location = request.POST.get("location", "").strip()
            meeting.chair_id = request.POST.get("chair") or meeting.chair_id
            meeting.recorder_id = request.POST.get("recorder") or meeting.recorder_id
            meeting.agenda = request.POST.get("agenda", "").strip()
            meeting.save()
            attendee_ids = request.POST.getlist("attendees")
            meeting.attendees.set(attendee_ids)
            absentee_ids = request.POST.getlist("absentees")
            meeting.absentees.set(absentee_ids)
            messages.success(request, "基本資訊已更新")
        elif action == "save_item":
            item = get_object_or_404(meeting.items, pk=request.POST.get("item_id"))
            item.topic = request.POST.get("topic", item.topic).strip()
            item.progress_note = request.POST.get("progress_note", "").strip()
            item.issue = request.POST.get("issue", "").strip()
            item.decision = request.POST.get("decision", "").strip()
            item.next_action = request.POST.get("next_action", "").strip()
            item.owner_id = request.POST.get("owner") or None
            raw_due = request.POST.get("due_date", "")
            item.due_date = date.fromisoformat(raw_due) if raw_due else None
            item.carry_to_next = request.POST.get("carry_to_next") == "on"
            item.save()
            messages.success(request, f"「{item.topic}」已更新。")
        
        elif action == "delete_item":
            item = get_object_or_404(meeting.items, pk=request.POST.get("item_id"))
            item.delete()
            messages.success(request, "已刪除議案項目")
        return redirect("meeting_detail", pk=meeting.pk)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["active_users"] = User.objects.filter(is_active=True).order_by("display_name")
        context["can_edit"] = self.request.user.role != User.Role.VIEWER and self.object.status != Meeting.Status.COMPLETED
        return context


class MeetingEndView(LoginRequiredMixin, View):
    def post(self, request, pk):
        if request.user.role == User.Role.VIEWER:
            raise PermissionDenied
        meeting = get_object_or_404(Meeting, pk=pk)
        meeting.status = Meeting.Status.COMPLETED
        meeting.ended_at = timezone.now()
        meeting.save(update_fields=["status", "ended_at"])
        messages.success(request, "會議已結束，紀錄已保存；勾選延續追蹤的事項會在下次會議帶入。")
        return redirect("meeting_center")


class TaskListView(TaskQuerysetMixin, ListView):
    template_name = "workitems/task_list.html"
    context_object_name = "tasks"
    paginate_by = 20

    def get_available_sections(self):
        sections = Department.objects.filter(
            is_active=True,
            unit_type__in=[Department.UnitType.DEPARTMENT, Department.UnitType.SECTION, Department.UnitType.GROUP],
        ).select_related("parent", "parent__parent").order_by("code")
        user = self.request.user
        if not (user.is_superuser or user.role == User.Role.ADMIN):
            scope_ids = department_scope_ids(user.department)
            sections = sections.filter(pk__in=scope_ids)
        units = list(sections)
        unit_ids = {unit.pk for unit in units}
        ordered = []

        def append_children(parent_id):
            for unit in units:
                if unit.parent_id == parent_id and unit.pk not in {item.pk for item in ordered}:
                    ordered.append(unit)
                    append_children(unit.pk)

        for unit in units:
            if unit.unit_type == Department.UnitType.DEPARTMENT or unit.parent_id not in unit_ids:
                if unit not in ordered:
                    ordered.append(unit)
                    append_children(unit.pk)
        for unit in units:
            if unit not in ordered:
                ordered.append(unit)
        return ordered

    def get_queryset(self):
        queryset = super().get_queryset()
        available_sections = self.get_available_sections()
        allowed_ids = {section.pk for section in available_sections}
        organization_filter_submitted = (
            "organization_filter_submitted" in self.request.GET
            or "departments" in self.request.GET
        )
        if organization_filter_submitted:
            selected_values = self.request.GET.getlist("departments")
        else:
            selected_values = [
                str(pk)
                for pk in self.request.user.preferred_task_departments.values_list("pk", flat=True)
            ]
        self.selected_section_ids = {int(value) for value in selected_values if value.isdigit() and int(value) in allowed_ids}
        if organization_filter_submitted:
            self.request.user.preferred_task_departments.set(self.selected_section_ids)
        focus_values = self.request.GET.getlist("focus_departments")
        legacy_focus_value = self.request.GET.get("focus_department", "")
        if not focus_values and legacy_focus_value:
            focus_values = [legacy_focus_value]
        self.focus_department_ids = {
            int(value)
            for value in focus_values
            if value.isdigit() and int(value) in self.selected_section_ids
        }
        if self.selected_section_ids:
            filter_ids = set()
            for unit in available_sections:
                if unit.pk in (self.focus_department_ids or self.selected_section_ids):
                    filter_ids.update(department_scope_ids(unit))
            queryset = queryset.filter(owner_department_id__in=filter_ids)
        q = self.request.GET.get("q", "").strip()
        if q:
            queryset = queryset.filter(
                Q(task_no__icontains=q) | 
                Q(title__icontains=q) | 
                Q(description__icontains=q) |
                Q(primary_assignee__display_name__icontains=q) |
                Q(requestor__display_name__icontains=q)
            )
        for field in ("status", "priority", "category", "owner_department", "primary_assignee"):
            value = self.request.GET.get(field)
            if value:
                queryset = queryset.filter(**{field if field in {"status", "priority"} else f"{field}_id": value})
        if not self.request.GET.get("status"):
            queryset = queryset.exclude(status__in=[Task.Status.COMPLETED, Task.Status.CANCELLED])
        
        sort = self.request.GET.get("sort")
        order = self.request.GET.get("order", "asc")
        valid_sort_fields = {
            "task_no": "task_no",
            "status": "status",
            "request_department": "request_department__name",
            "owner_department": "owner_department__name",
            "priority": "priority",
            "progress": "progress",
            "due_date": "due_date"
        }
        if sort in valid_sort_fields:
            sort_field = valid_sort_fields[sort]
            if order == "desc":
                sort_field = "-" + sort_field
            return queryset.order_by(sort_field)
            
        return queryset.order_by("due_date", "-priority", "-created_at")

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["current_sort"] = self.request.GET.get("sort", "")
        context["current_order"] = self.request.GET.get("order", "asc")
        for task in context["tasks"]:
            task.user_can_edit = can_edit_task(self.request.user, task)
        sections = self.get_available_sections()
        visible_open_tasks = visible_tasks(
            self.request.user,
            Task.objects.exclude(status__in=[Task.Status.COMPLETED, Task.Status.CANCELLED]),
        )
        visible_all_tasks = visible_tasks(
            self.request.user,
            Task.objects.exclude(status=Task.Status.CANCELLED),
        )
        selected_unit_summaries = []
        color_index = 0
        for section in sections:
            section.selected = section.pk in getattr(self, "selected_section_ids", set())
            section.depth = {
                Department.UnitType.DEPARTMENT: 0,
                Department.UnitType.SECTION: 1,
                Department.UnitType.GROUP: 2,
            }[section.unit_type]
            count_ids = department_scope_ids(section)
            scoped_tasks = visible_open_tasks.filter(owner_department_id__in=count_ids)
            scoped_all_tasks = visible_all_tasks.filter(owner_department_id__in=count_ids)
            section.open_count = scoped_tasks.count()
            if section.selected:
                overdue_filter = Q(due_date__lt=timezone.localdate())
                total_count = scoped_all_tasks.count()
                completed_count = scoped_all_tasks.filter(status=Task.Status.COMPLETED).count()
                focus_query = self.request.GET.copy()
                focus_query["organization_filter_submitted"] = "1"
                focus_query.setlist("departments", [str(pk) for pk in self.selected_section_ids])
                toggled_focus_ids = set(self.focus_department_ids)
                if section.pk in toggled_focus_ids:
                    toggled_focus_ids.remove(section.pk)
                else:
                    toggled_focus_ids.add(section.pk)
                focus_query.pop("focus_department", None)
                focus_query.pop("focus_departments", None)
                if toggled_focus_ids:
                    focus_query.setlist("focus_departments", [str(pk) for pk in sorted(toggled_focus_ids)])
                focus_query.pop("page", None)
                selected_unit_summaries.append({
                    "unit": section,
                    "total": total_count,
                    "open_total": section.open_count,
                    "completed": completed_count,
                    "completion_rate": round(completed_count / total_count * 100) if total_count else 0,
                    "high_priority": scoped_tasks.filter(priority__in=[Task.Priority.HIGH, Task.Priority.URGENT]).count(),
                    "pending": scoped_tasks.filter(status__in=[Task.Status.NOT_STARTED, Task.Status.WAITING]).count(),
                    "in_progress": scoped_tasks.filter(status=Task.Status.IN_PROGRESS).count(),
                    "overdue": scoped_tasks.filter(overdue_filter).count(),
                    "focus_url": f"?{focus_query.urlencode()}",
                    "is_focused": section.pk in self.focus_department_ids,
                    "color_index": color_index % 6,
                })
                color_index += 1
        context.update({
            "statuses": Task.Status.choices,
            "priorities": Task.Priority.choices,
            "categories": TaskCategory.objects.filter(is_active=True),
            "query": self.request.GET,
            "can_create": self.request.user.role != User.Role.VIEWER,
            "section_departments": sections,
            "selected_section_count": len(getattr(self, "selected_section_ids", set())),
            "selected_unit_summaries": selected_unit_summaries,
            "focus_department_ids": self.focus_department_ids,
            "export_fields": EXPORT_FIELDS,
        })
        return context


class TaskHistoryView(TaskQuerysetMixin, ListView):
    template_name = "workitems/task_history.html"
    context_object_name = "tasks"
    paginate_by = 30

    def get_queryset(self):
        queryset = super().get_queryset().filter(status__in=[Task.Status.COMPLETED, Task.Status.CANCELLED])
        query = self.request.GET.get("q", "").strip()
        if query:
            queryset = queryset.filter(Q(task_no__icontains=query) | Q(title__icontains=query) | Q(description__icontains=query))
        return queryset.order_by("-completed_at", "-updated_at")

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        query = self.request.GET.get("q", "").strip()
        completed_notes = IssueNote.objects.filter(status=IssueNote.Status.DONE).select_related(
            "created_by", "resolved_by"
        ).prefetch_related("attachments")
        if query:
            completed_notes = completed_notes.filter(
                Q(title__icontains=query) | Q(content__icontains=query) | Q(reported_by_name__icontains=query)
            )
        completed_notes = completed_notes.order_by("-resolved_at", "-updated_at")
        for note in completed_notes:
            note.user_can_toggle = (
                self.request.user.is_superuser
                or self.request.user.role in {User.Role.ADMIN, User.Role.MANAGER}
                or note.created_by_id == self.request.user.id
            ) and self.request.user.role != User.Role.VIEWER
        context["completed_notes"] = completed_notes
        context["completed_note_count"] = completed_notes.count()
        return context


class TaskCalendarView(TaskQuerysetMixin, TemplateView):
    template_name = "workitems/task_calendar.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        today = timezone.localdate()
        try:
            year = int(self.request.GET.get("year", today.year))
            month = int(self.request.GET.get("month", today.month))
            if not 1 <= month <= 12 or not 2000 <= year <= 2100:
                raise ValueError
        except (TypeError, ValueError):
            year, month = today.year, today.month

        first_day = date(year, month, 1)
        next_month = date(year + 1, 1, 1) if month == 12 else date(year, month + 1, 1)
        previous_month = date(year - 1, 12, 1) if month == 1 else date(year, month - 1, 1)
        tasks = self.get_queryset().exclude(status=Task.Status.CANCELLED).filter(
            due_date__gte=first_day, due_date__lt=next_month,
        ).order_by("due_date", "-priority", "title")
        tasks_by_date = {}
        for task in tasks:
            task.user_can_edit = can_edit_task(self.request.user, task)
            tasks_by_date.setdefault(task.due_date, []).append(task)

        weeks = []
        for week in calendar.Calendar(firstweekday=6).monthdatescalendar(year, month):
            weeks.append([{
                "date": day,
                "in_month": day.month == month,
                "is_today": day == today,
                "tasks": tasks_by_date.get(day, []),
            } for day in week])

        context.update({
            "calendar_weeks": weeks,
            "month_label": f"{year} 年 {month} 月",
            "previous_month": previous_month,
            "next_month": next_month,
            "task_count": tasks.count(),
        })
        return context


class TaskExcelExportView(LoginRequiredMixin, View):
    def post(self, request):
        mode = request.POST.get("export_mode", "filtered")
        include_history = request.POST.get("include_history") == "on"
        split_by_status = request.POST.get("split_by_status") == "on"
        queryset = visible_tasks(
            request.user,
            Task.objects.select_related(
                "project", "category", "request_department", "requestor", "owner_department", "primary_assignee"
            ).prefetch_related("tags", "collaborators"),
        )
        if not include_history:
            queryset = queryset.exclude(status__in=[Task.Status.COMPLETED, Task.Status.CANCELLED])

        mode_labels = {"selected": "勾選項目", "filtered": "目前篩選結果", "all": "全部授權資料"}
        if mode == "selected":
            selected_ids = {
                int(value) for value in request.POST.get("selected_ids", "").split(",") if value.strip().isdigit()
            }
            queryset = queryset.filter(pk__in=selected_ids)
        elif mode == "filtered":
            filters = QueryDict(request.POST.get("filter_query", ""))
            available_sections = TaskListView()
            available_sections.request = request
            sections = available_sections.get_available_sections()
            allowed_ids = {section.pk for section in sections}
            if "organization_filter_submitted" in filters or "departments" in filters:
                department_values = filters.getlist("departments")
            else:
                department_values = [str(pk) for pk in request.user.preferred_task_departments.values_list("pk", flat=True)]
            selected_departments = {
                int(value) for value in department_values if value.isdigit() and int(value) in allowed_ids
            }
            focused_values = filters.getlist("focus_departments")
            focused_departments = {
                int(value) for value in focused_values if value.isdigit() and int(value) in selected_departments
            }
            if selected_departments:
                scope_ids = set()
                for unit in sections:
                    if unit.pk in (focused_departments or selected_departments):
                        scope_ids.update(department_scope_ids(unit))
                queryset = queryset.filter(owner_department_id__in=scope_ids)
            query = filters.get("q", "").strip()
            if query:
                queryset = queryset.filter(Q(task_no__icontains=query) | Q(title__icontains=query) | Q(description__icontains=query))
            for field in ("status", "priority", "category", "owner_department", "primary_assignee"):
                value = filters.get(field)
                if value:
                    queryset = queryset.filter(**{field if field in {"status", "priority"} else f"{field}_id": value})
        elif mode != "all":
            return JsonResponse({"error": "不支援的匯出模式。"}, status=400)

        selected_fields = request.POST.getlist("fields")
        workbook = build_task_workbook(
            queryset.order_by("due_date", "task_no"),
            selected_fields,
            split_by_status,
            mode_labels[mode],
            request.build_absolute_uri("/").rstrip("/"),
        )
        response = HttpResponse(
            workbook.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        filename = f"TaskFlow_work_export_{timezone.localdate():%Y%m%d}.xlsx"
        response["Content-Disposition"] = f"attachment; filename=\"{filename}\""
        return response


class SystemSettingsView(LoginRequiredMixin, TemplateView):
    template_name = "workitems/system_settings.html"

    def dispatch(self, request, *args, **kwargs):
        if not request.user.is_staff:
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["task_categories"] = TaskCategory.objects.annotate(usage_count=Count("task"))
        context["task_tags"] = Tag.objects.annotate(usage_count=Count("tasktag"))
        return context

    def post(self, request, *args, **kwargs):
        action = request.POST.get("action")
        try:
            if action == "save_category":
                category_id = request.POST.get("id")
                category = get_object_or_404(TaskCategory, pk=category_id) if category_id else TaskCategory()
                category.name = request.POST.get("name", "").strip()
                category.sort_order = int(request.POST.get("sort_order") or 0)
                category.is_active = request.POST.get("is_active") == "on"
                category.full_clean()
                category.save()
                messages.success(request, "工作類別已儲存。")
                return redirect(f"{reverse_lazy('system_settings')}#categories")
            if action == "delete_category":
                category = get_object_or_404(TaskCategory, pk=request.POST.get("id"))
                usage_count = category.task_set.count()
                if usage_count:
                    category.is_active = False
                    category.save(update_fields=["is_active"])
                    messages.warning(request, f"「{category.name}」已被 {usage_count} 件工作使用，已改為停用並保留歷史資料。")
                else:
                    name = category.name
                    category.delete()
                    messages.success(request, f"工作類別「{name}」已刪除。")
                return redirect(f"{reverse_lazy('system_settings')}#categories")
            if action == "save_tag":
                tag_id = request.POST.get("id")
                tag = get_object_or_404(Tag, pk=tag_id) if tag_id else Tag()
                tag.name = request.POST.get("name", "").strip()
                color = request.POST.get("color", "").strip().upper()
                if len(color) != 7 or not color.startswith("#") or any(character not in "0123456789ABCDEF" for character in color[1:]):
                    raise ValidationError("顏色請使用 #RRGGBB 格式。")
                tag.color = color
                tag.is_active = request.POST.get("is_active") == "on"
                tag.full_clean()
                tag.save()
                messages.success(request, "工作標籤已儲存。")
                return redirect(f"{reverse_lazy('system_settings')}#tags")
            if action == "delete_tag":
                tag = get_object_or_404(Tag, pk=request.POST.get("id"))
                usage_count = tag.tasktag_set.count()
                if usage_count:
                    tag.is_active = False
                    tag.save(update_fields=["is_active"])
                    messages.warning(request, f"「{tag.name}」已被 {usage_count} 件工作使用，已改為停用並保留歷史資料。")
                else:
                    name = tag.name
                    tag.delete()
                    messages.success(request, f"工作標籤「{name}」已刪除。")
                return redirect(f"{reverse_lazy('system_settings')}#tags")
        except (ValidationError, ValueError) as error:
            detail = error.message_dict if hasattr(error, "message_dict") else error.messages if hasattr(error, "messages") else str(error)
            messages.error(request, f"儲存失敗：{detail}")
        return redirect("system_settings")


class IssueNoteBoardView(LoginRequiredMixin, ListView):
    model = IssueNote
    template_name = "workitems/issue_note_board.html"
    context_object_name = "notes"

    def get_queryset(self):
        queryset = IssueNote.objects.exclude(status=IssueNote.Status.DONE).select_related(
            "created_by", "resolved_by"
        ).prefetch_related("attachments")
        status = self.request.GET.get("status")
        if status in IssueNote.Status.values:
            queryset = queryset.filter(status=status)
        return queryset

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        notes = list(context["notes"])
        for note in notes:
            note.user_can_toggle = (
                self.request.user.is_superuser
                or self.request.user.role in {User.Role.ADMIN, User.Role.MANAGER}
                or note.created_by_id == self.request.user.id
            ) and self.request.user.role != User.Role.VIEWER
        context["form"] = kwargs.get("form") or IssueNoteForm(initial={"reported_by_name": self.request.user.display_name})
        context["can_create"] = self.request.user.role != User.Role.VIEWER
        context["open_count"] = IssueNote.objects.exclude(status=IssueNote.Status.DONE).count()
        icons = {IssueNote.Status.INBOX: "📥", IssueNote.Status.MUST_DO: "🔥", IssueNote.Status.IN_PROGRESS: "⏳"}
        context["note_columns"] = [
            {"value": value, "label": label, "icon": icons[value], "notes": [note for note in notes if note.status == value]}
            for value, label in IssueNote.Status.choices if value != IssueNote.Status.DONE
        ]
        return context

    def post(self, request, *args, **kwargs):
        if request.user.role == User.Role.VIEWER:
            raise PermissionDenied
        form = IssueNoteForm(request.POST)
        attachments = request.FILES.getlist("attachments")
        allowed_extensions = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".pdf", ".doc", ".docx", ".xls", ".xlsx", ".ppt", ".pptx", ".txt", ".csv"}
        for attachment in attachments:
            extension = Path(attachment.name).suffix.lower()
            if extension not in allowed_extensions:
                form.add_error(None, f"不支援附件格式：{attachment.name}")
            if attachment.size > 10 * 1024 * 1024:
                form.add_error(None, f"附件超過 10 MB：{attachment.name}")
        if form.is_valid():
            note = form.save(commit=False)
            note.created_by = request.user
            note.save()
            for attachment in attachments:
                IssueNoteAttachment.objects.create(
                    note=note,
                    file=attachment,
                    original_name=Path(attachment.name).name,
                    file_size=attachment.size,
                    uploaded_by=request.user,
                )
            messages.success(request, "問題便簽已記錄。")
            return redirect("issue_note_board")
        self.object_list = self.get_queryset()
        return self.render_to_response(self.get_context_data(form=form))


class IssueNoteAttachmentDownloadView(LoginRequiredMixin, View):
    def get(self, request, pk):
        attachment = get_object_or_404(IssueNoteAttachment, pk=pk)
        return FileResponse(attachment.file.open("rb"), as_attachment=True, filename=attachment.original_name)


class IssueNoteToggleView(LoginRequiredMixin, View):
    NEXT_STATUS = {
        IssueNote.Status.INBOX: IssueNote.Status.MUST_DO,
        IssueNote.Status.MUST_DO: IssueNote.Status.IN_PROGRESS,
        IssueNote.Status.IN_PROGRESS: IssueNote.Status.DONE,
        IssueNote.Status.DONE: IssueNote.Status.INBOX,
    }

    def post(self, request, pk):
        note = get_object_or_404(IssueNote, pk=pk)
        allowed = request.user.is_superuser or request.user.role in {User.Role.ADMIN, User.Role.MANAGER} or note.created_by_id == request.user.id
        if not allowed or request.user.role == User.Role.VIEWER:
            raise PermissionDenied
        note.status = self.NEXT_STATUS.get(note.status, IssueNote.Status.INBOX)
        if note.status == IssueNote.Status.DONE:
            note.resolved_by = request.user
            note.resolved_at = timezone.now()
        else:
            note.resolved_by = None
            note.resolved_at = None
        note.save(update_fields=["status", "resolved_by", "resolved_at", "updated_at"])
        return redirect("issue_note_board")


class TaskQuickStatusView(LoginRequiredMixin, View):
    NEXT_STATUS = {
        Task.Status.NOT_STARTED: Task.Status.IN_PROGRESS,
        Task.Status.IN_PROGRESS: Task.Status.COMPLETED,
        Task.Status.WAITING: Task.Status.IN_PROGRESS,
    }

    def post(self, request, pk):
        task = get_object_or_404(Task, pk=pk)
        if not can_edit_task(request.user, task):
            raise PermissionDenied
        requested_status = request.POST.get("target_status", "")
        allowed_statuses = {
            Task.Status.NOT_STARTED,
            Task.Status.IN_PROGRESS,
            Task.Status.WAITING,
            Task.Status.COMPLETED,
        }
        next_status = requested_status if requested_status in allowed_statuses else self.NEXT_STATUS.get(task.status)
        if not next_status:
            return JsonResponse({"error": "此狀態無法快速切換。"}, status=400)
        if next_status == Task.Status.WAITING:
            task.waiting_reason = request.POST.get("waiting_reason", "").strip() or task.waiting_reason.strip()
        elif task.status == Task.Status.WAITING:
            task.waiting_reason = ""
        task.status = next_status
        if next_status == Task.Status.IN_PROGRESS and task.progress == 0:
            task.progress = 1
        elif next_status == Task.Status.NOT_STARTED:
            task.progress = 0
        task.updated_by = request.user
        task.save()
        return JsonResponse({
            "status": task.status,
            "label": task.get_status_display(),
            "progress": task.progress,
            "waiting_reason": task.waiting_reason,
            "moved_to_history": task.status == Task.Status.COMPLETED,
            "project": ({
                "progress": task.project.progress,
                "task_count": task.project.tasks.count(),
                "completed_count": task.project.tasks.filter(status=Task.Status.COMPLETED).count(),
            } if task.project_id else None),
        })


class ProjectBoardView(LoginRequiredMixin, ListView):
    model = Project
    template_name = "workitems/project_board.html"
    context_object_name = "projects"

    def get_queryset(self):
        queryset = Project.objects.select_related("owner_department", "manager").annotate(
            task_count=Count("tasks", distinct=True),
            completed_count=Count("tasks", filter=Q(tasks__status=Task.Status.COMPLETED), distinct=True),
        )
        if not (self.request.user.is_superuser or self.request.user.role == User.Role.ADMIN):
            visible_ids = visible_tasks(self.request.user, Task.objects.all()).values("project_id")
            queryset = queryset.filter(Q(created_by=self.request.user) | Q(manager=self.request.user) | Q(id__in=visible_ids)).distinct()
        return queryset

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        projects = list(context["projects"])
        selected_id = self.request.GET.get("project")
        selected = next((item for item in projects if str(item.pk) == selected_id), projects[0] if projects else None)
        task_queryset = visible_tasks(self.request.user, Task.objects.all()).filter(project=selected).exclude(
            status__in=[Task.Status.COMPLETED, Task.Status.CANCELLED]
        ).select_related(
            "request_department", "owner_department", "primary_assignee", "category"
        ) if selected else Task.objects.none()
        tasks = list(task_queryset)
        for task in tasks:
            task.user_can_edit = can_edit_task(self.request.user, task)
        assignees = User.objects.filter(is_active=True, department=selected.owner_department).order_by("display_name") if selected else User.objects.none()
        context.update({
            "projects": projects,
            "selected_project": selected,
            "project_tasks": tasks,
            "can_create": self.request.user.role != User.Role.VIEWER,
            "quick_assignees": assignees,
            "quick_rows": range(6),
            "priorities": Task.Priority.choices,
        })
        return context


class ProjectBulkTaskCreateView(LoginRequiredMixin, View):
    def post(self, request, pk):
        project = get_object_or_404(Project.objects.select_related("owner_department", "manager"), pk=pk)
        board_url = f"{reverse_lazy('project_board')}?project={project.pk}#project-task-list"
        user = request.user
        allowed = user.is_superuser or user.role == User.Role.ADMIN or project.manager_id == user.id or project.created_by_id == user.id
        if user.role == User.Role.MANAGER and user.department:
            allowed = allowed or project.owner_department_id in department_scope_ids(user.department)
        if user.role == User.Role.VIEWER or not allowed:
            raise PermissionDenied

        titles = request.POST.getlist("title")
        descriptions = request.POST.getlist("description")
        assignees = request.POST.getlist("assignee")
        due_dates = request.POST.getlist("due_date")
        priorities = request.POST.getlist("priority")
        category = TaskCategory.objects.filter(name="一般工作", is_active=True).first() or TaskCategory.objects.filter(is_active=True).first()
        if not category:
            messages.error(request, "請先建立至少一個工作類別。")
            return redirect(board_url)

        created = 0
        try:
            with transaction.atomic():
                for index, raw_title in enumerate(titles):
                    title = raw_title.strip()
                    if not title:
                        continue
                    description = descriptions[index].strip() if index < len(descriptions) else ""
                    assignee_id = assignees[index] if index < len(assignees) else ""
                    assignee = User.objects.filter(pk=assignee_id, is_active=True, department=project.owner_department).first() or project.manager
                    raw_due = due_dates[index] if index < len(due_dates) else ""
                    due = date.fromisoformat(raw_due) if raw_due else project.due_date
                    priority = priorities[index] if index < len(priorities) and priorities[index] in Task.Priority.values else Task.Priority.MEDIUM
                    Task.objects.create(
                        project=project,
                        title=title,
                        description=description or "由專案看板快速建立，可進入工作單補充詳細說明。",
                        category=category,
                        request_department=project.owner_department,
                        requestor=project.manager,
                        owner_department=project.owner_department,
                        primary_assignee=assignee,
                        priority=priority,
                        status=Task.Status.NOT_STARTED,
                        progress=0,
                        planned_start_date=project.start_date,
                        due_date=due,
                        created_by=user,
                        updated_by=user,
                    )
                    created += 1
        except (ValidationError, ValueError, TypeError) as error:
            messages.error(request, f"快速建立失敗：{error}")
            return redirect(board_url)
        messages.success(request, f"已新增 {created} 件專案工作。" if created else "沒有可新增的工作名稱。")
        return redirect(board_url)


class ProjectCreateView(LoginRequiredMixin, CreateView):
    model = Project
    form_class = ProjectForm
    template_name = "workitems/project_form.html"

    def dispatch(self, request, *args, **kwargs):
        if request.user.is_authenticated and request.user.role == User.Role.VIEWER:
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def form_valid(self, form):
        form.instance.created_by = self.request.user
        return super().form_valid(form)

    def get_success_url(self):
        return reverse_lazy("project_board") + f"?project={self.object.pk}"


class TaskCreateView(LoginRequiredMixin, CreateView):
    model = Task
    form_class = TaskForm
    template_name = "workitems/task_form.html"

    def dispatch(self, request, *args, **kwargs):
        if not request.user.is_authenticated:
            return self.handle_no_permission()
        if request.user.role == User.Role.VIEWER:
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def form_valid(self, form):
        form.instance.created_by = self.request.user
        form.instance.updated_by = self.request.user
        return super().form_valid(form)

    def get_initial(self):
        initial = super().get_initial()
        if self.request.GET.get("project"):
            initial["project"] = self.request.GET["project"]
        return initial

    def get_success_url(self):
        return reverse_lazy("task_detail", kwargs={"pk": self.object.pk})


class TaskDetailView(TaskQuerysetMixin, UpdateView):
    form_class = TaskForm
    template_name = "workitems/task_detail.html"
    context_object_name = "task"

    def post(self, request, *args, **kwargs):
        self.object = self.get_object()
        if not can_edit_task(request.user, self.object):
            raise PermissionDenied
        return super().post(request, *args, **kwargs)

    def form_valid(self, form):
        form.instance.updated_by = self.request.user
        messages.success(self.request, "工作內容已更新。")
        return super().form_valid(form)

    def get_success_url(self):
        return reverse_lazy("task_detail", kwargs={"pk": self.object.pk})

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["can_edit"] = can_edit_task(self.request.user, self.object)
        return context


class TaskUpdateView(TaskQuerysetMixin, UpdateView):
    form_class = TaskForm
    template_name = "workitems/task_form.html"

    def dispatch(self, request, *args, **kwargs):
        if not request.user.is_authenticated:
            return self.handle_no_permission()
        task = self.get_object()
        if not can_edit_task(request.user, task):
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def form_valid(self, form):
        form.instance.updated_by = self.request.user
        return super().form_valid(form)

    def get_success_url(self):
        return reverse_lazy("task_detail", kwargs={"pk": self.object.pk})


class DepartmentUsersView(LoginRequiredMixin, View):
    def get(self, request, pk):
        users = User.objects.filter(is_active=True, department_id=pk).values("id", "display_name", "username")
        return JsonResponse({"users": list(users)})


class UserSearchView(LoginRequiredMixin, View):
    def get(self, request):
        query = request.GET.get("q", "").strip()
        users = User.objects.filter(is_active=True, department__isnull=False).select_related("department")
        if query:
            users = users.filter(Q(display_name__icontains=query) | Q(username__icontains=query))
        users = users.order_by("display_name", "username")[:20]
        return JsonResponse({"users": [{
            "id": user.pk,
            "display_name": user.display_name,
            "username": user.username,
            "department_id": user.department_id,
            "department_name": user.department.name,
        } for user in users]})


import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side
from django.http import HttpResponse

class MeetingExcelExportView(LoginRequiredMixin, View):
    def get(self, request, pk):
        meeting = get_object_or_404(Meeting.objects.prefetch_related("items__owner", "attendees", "absentees"), pk=pk)
        
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "會議紀錄"
        
        # Styles
        title_font = Font(size=16, bold=True)
        bold_font = Font(bold=True)
        border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
        align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
        align_left = Alignment(horizontal='left', vertical='top', wrap_text=True)
        
        # Title
        ws.append([f"{meeting.title} 會議記錄"])
        ws.merge_cells('A1:E1')
        ws['A1'].font = title_font
        ws['A1'].alignment = align_center
        ws.append([])
        
        # Meta
        date_str = meeting.started_at.strftime('%Y/%m/%d') if meeting.started_at else (meeting.scheduled_at.strftime('%Y/%m/%d') if meeting.scheduled_at else '')
        loc_str = meeting.location or (meeting.series.location if meeting.series else '') or '(依行事曆通知)'
        chair_str = str(meeting.chair) if meeting.chair else ''
        rec_str = str(meeting.recorder) if meeting.recorder else ''
        
        attendees = "、".join([user.display_name for user in meeting.attendees.all() if user.display_name != "系統管理員"]) or "無"
        absentees = "、".join([user.display_name for user in meeting.absentees.all() if user.display_name != "系統管理員"]) or "無"
        
        ws.append(["會議名稱：", meeting.title])
        ws.append(["會議日期：", date_str])
        ws.append(["會議地點：", loc_str])
        ws.append(["主席：", chair_str, "記錄：", rec_str])
        ws.append(["出席人員：", attendees])
        ws.append(["缺席人員：", absentees])
        ws.append([])
        
        # Table Header
        headers = ["編號", "議案", "決議", "執行人", "完成期限"]
        ws.append(headers)
        for col_num, cell in enumerate(ws[ws.max_row], 1):
            cell.font = bold_font
            cell.border = border
            cell.alignment = align_center
            
        # Table Body
        for idx, item in enumerate(meeting.items.all(), 1):
            owner_str = str(item.owner) if item.owner else "All"
            due_str = item.due_date.strftime('%Y-%m-%d') if item.due_date else "-"
            row = [
                str(idx),
                item.topic,
                item.decision,
                owner_str,
                due_str
            ]
            ws.append(row)
            for col_num, cell in enumerate(ws[ws.max_row], 1):
                cell.border = border
                cell.alignment = align_left if col_num in (2, 3) else align_center
                
        # Footer
        ws.append([])
        ws.append(["※未出席會議的人員如對會議內容有異議，請於3個工作天內提出，並通知所有會議相關人員"])
        ws['A' + str(ws.max_row)].font = Font(color="FF0000")
        
        # Column Widths
        ws.column_dimensions['A'].width = 8
        ws.column_dimensions['B'].width = 30
        ws.column_dimensions['C'].width = 50
        ws.column_dimensions['D'].width = 15
        ws.column_dimensions['E'].width = 15
        
        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = f'attachment; filename="Meeting_Record_{date_str.replace("/","")}.xlsx"'
        wb.save(response)
        return response


class MeetingSeriesUpdateView(LoginRequiredMixin, UpdateView):
    model = MeetingSeries
    form_class = MeetingSeriesForm
    template_name = "workitems/meeting_series_form.html"
    success_url = reverse_lazy("meeting_center")

    def dispatch(self, request, *args, **kwargs):
        if request.user.is_authenticated and request.user.role == User.Role.VIEWER:
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def form_valid(self, form):
        messages.success(self.request, "會議系列已更新")
        return super().form_valid(form)


class MeetingSeriesDeleteView(LoginRequiredMixin, View):
    def post(self, request, pk):
        if request.user.role == User.Role.VIEWER:
            raise PermissionDenied
        series = get_object_or_404(MeetingSeries, pk=pk)
        name = series.name
        series.delete()
        messages.success(request, f"已刪除會議系列「{name}」")
        return redirect("meeting_center")
