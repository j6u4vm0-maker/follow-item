from io import BytesIO

from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


EXPORT_FIELDS = {
    "task_no": ("工作編號", lambda task: task.task_no),
    "title": ("工作名稱", lambda task: task.title),
    "description": ("工作說明", lambda task: task.description),
    "project": ("所屬專案", lambda task: task.project.name if task.project else ""),
    "category": ("工作類別", lambda task: task.category.name),
    "tags": ("標籤", lambda task: "、".join(tag.name for tag in task.tags.all())),
    "status": ("狀態", lambda task: task.get_status_display()),
    "priority": ("優先度", lambda task: task.get_priority_display()),
    "progress": ("進度", lambda task: task.progress / 100),
    "request_department": ("需求單位", lambda task: task.request_department.name),
    "requestor": ("需求人", lambda task: str(task.requestor)),
    "owner_department": ("負責單位", lambda task: task.owner_department.name),
    "primary_assignee": ("主要負責人", lambda task: str(task.primary_assignee)),
    "collaborators": ("協作人員", lambda task: "、".join(str(user) for user in task.collaborators.all())),
    "request_date": ("提出日期", lambda task: task.request_date),
    "planned_start_date": ("預計開始日", lambda task: task.planned_start_date),
    "due_date": ("預計完成日", lambda task: task.due_date),
    "completed_at": ("完成時間", lambda task: timezone.localtime(task.completed_at).replace(tzinfo=None) if task.completed_at else None),
    "updated_at": ("最後更新", lambda task: timezone.localtime(task.updated_at).replace(tzinfo=None)),
}


def build_task_workbook(tasks, selected_fields, split_by_status, mode_label, base_url):
    tasks = list(tasks)
    fields = [field for field in selected_fields if field in EXPORT_FIELDS]
    if not fields:
        fields = list(EXPORT_FIELDS)
    workbook = Workbook()
    workbook.remove(workbook.active)
    groups = [("工作清單", tasks)]
    if split_by_status:
        status_order = ["not_started", "in_progress", "waiting", "completed", "cancelled"]
        status_labels = dict(tasks[0].Status.choices) if tasks else {
            "not_started": "未開始", "in_progress": "進行中", "waiting": "等待", "completed": "完成", "cancelled": "取消",
        }
        groups = [(status_labels[value], [task for task in tasks if task.status == value]) for value in status_order]
        groups = [(name, rows) for name, rows in groups if rows] or [("工作清單", [])]

    for sheet_name, rows in groups:
        sheet = workbook.create_sheet(sheet_name[:31])
        sheet.sheet_view.showGridLines = False
        last_column = max(1, len(fields))
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_column)
        sheet.cell(1, 1, f"TaskFlow 工作匯出｜{sheet_name}")
        sheet.cell(1, 1).font = Font(name="Arial", size=15, bold=True, color="16324F")
        sheet.cell(2, 1, f"匯出範圍：{mode_label}")
        sheet.cell(2, 1).font = Font(name="Arial", size=10, italic=True, color="6B7788")
        sheet.cell(3, 1, f"資料筆數：{len(rows)}　匯出時間：{timezone.localtime():%Y/%m/%d %H:%M}")
        sheet.cell(3, 1).font = Font(name="Arial", size=10, color="6B7788")
        for column, field in enumerate(fields, 1):
            cell = sheet.cell(4, column, EXPORT_FIELDS[field][0])
            cell.fill = PatternFill("solid", fgColor="245B7A")
            cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        thin = Side(style="thin", color="DCE3EA")
        status_fills = {"未開始": "EEF1F5", "進行中": "DDEEFF", "等待": "FFF0CD", "完成": "DDF3E8", "取消": "E9EBEE"}
        for row_index, task in enumerate(rows, 5):
            for column, field in enumerate(fields, 1):
                value = EXPORT_FIELDS[field][1](task)
                cell = sheet.cell(row_index, column, value)
                cell.font = Font(name="Arial", size=10, color="172238")
                cell.alignment = Alignment(vertical="top", wrap_text=field in {"title", "description", "tags", "collaborators"})
                cell.border = Border(bottom=thin)
                if field in {"request_date", "planned_start_date", "due_date"} and value:
                    cell.number_format = "yyyy/mm/dd"
                elif field in {"completed_at", "updated_at"} and value:
                    cell.number_format = "yyyy/mm/dd hh:mm"
                elif field == "progress":
                    cell.number_format = "0%"
                elif field == "task_no":
                    cell.hyperlink = f"{base_url.rstrip('/')}{reverse('task_detail', args=[task.pk])}"
                    cell.font = Font(name="Arial", size=10, color="1677FF", underline="single")
                if field == "status" and value in status_fills:
                    cell.fill = PatternFill("solid", fgColor=status_fills[value])
            sheet.row_dimensions[row_index].height = 30
        sheet.freeze_panes = "A5"
        if rows:
            sheet.auto_filter.ref = f"A4:{get_column_letter(last_column)}{4 + len(rows)}"
        for column, field in enumerate(fields, 1):
            values = [str(EXPORT_FIELDS[field][0])] + [str(EXPORT_FIELDS[field][1](task) or "") for task in rows[:200]]
            width = min(max(len(value) for value in values) + 3, 48)
            if field == "description":
                width = min(max(width, 28), 48)
            sheet.column_dimensions[get_column_letter(column)].width = max(width, 12)
        sheet.row_dimensions[1].height = 24
        sheet.row_dimensions[4].height = 24

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
