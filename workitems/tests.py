from datetime import timedelta
from io import BytesIO

from django.core.exceptions import ValidationError
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from accounts.models import User
from organization.models import Department
from .models import IssueNote, Project, Tag, Task, TaskCategory
from .permissions import visible_tasks


class TaskTestBase(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.hq = Department.objects.create(code="HQ", name="總部", unit_type=Department.UnitType.HEADQUARTERS)
        cls.division = Department.objects.create(code="D1", name="營運處", unit_type=Department.UnitType.DIVISION, parent=cls.hq)
        cls.department = Department.objects.create(code="QA", name="品質部", unit_type=Department.UnitType.DEPARTMENT, parent=cls.division)
        cls.other_division = Department.objects.create(code="D2", name="研發處", unit_type=Department.UnitType.DIVISION, parent=cls.hq)
        cls.other_department = Department.objects.create(code="RD", name="研發部", unit_type=Department.UnitType.DEPARTMENT, parent=cls.other_division)
        cls.manager = User.objects.create_user("manager", password="Strong-pass-123", display_name="品質主管", role=User.Role.MANAGER, department=cls.department)
        cls.member = User.objects.create_user("member", password="Strong-pass-123", display_name="品質人員", role=User.Role.USER, department=cls.department)
        cls.outsider = User.objects.create_user("outsider", password="Strong-pass-123", display_name="研發人員", role=User.Role.USER, department=cls.other_department)
        cls.category = TaskCategory.objects.create(name="一般工作")

    def make_task(self, **overrides):
        values = {
            "title": "測試工作",
            "description": "工作說明",
            "category": self.category,
            "request_department": self.department,
            "requestor": self.member,
            "owner_department": self.department,
            "primary_assignee": self.member,
            "due_date": timezone.localdate() + timedelta(days=7),
            "created_by": self.manager,
            "updated_by": self.manager,
        }
        values.update(overrides)
        return Task.objects.create(**values)


class TaskModelTests(TaskTestBase):
    def test_generates_unique_yearly_task_numbers(self):
        first = self.make_task()
        second = self.make_task(title="第二件")
        self.assertRegex(first.task_no, r"^TASK-\d{4}-0001$")
        self.assertRegex(second.task_no, r"^TASK-\d{4}-0002$")

    def test_waiting_reason_is_optional(self):
        task = self.make_task(status=Task.Status.WAITING)
        self.assertEqual(task.waiting_reason, "")

    def test_completed_sets_progress_and_timestamp(self):
        task = self.make_task(status=Task.Status.COMPLETED, progress=40)
        self.assertEqual(task.progress, 100)
        self.assertIsNotNone(task.completed_at)

    def test_requestor_must_belong_to_request_department(self):
        with self.assertRaises(ValidationError):
            self.make_task(requestor=self.outsider)

    def test_project_progress_averages_its_tasks(self):
        project = Project.objects.create(name="年度專案", owner_department=self.department, manager=self.manager, start_date=timezone.localdate(), due_date=timezone.localdate() + timedelta(days=30), created_by=self.manager)
        self.make_task(project=project, progress=20)
        self.make_task(project=project, progress=80)
        self.assertEqual(project.progress, 50)


class TaskPermissionTests(TaskTestBase):
    def test_anonymous_create_redirects_to_login(self):
        response = self.client.get(reverse("task_create"))
        self.assertRedirects(response, f"{reverse('login')}?next={reverse('task_create')}")

    def test_task_form_uses_search_fields_for_requestor_and_primary_assignee(self):
        self.client.force_login(self.manager)
        response = self.client.get(reverse("task_create"))
        self.assertContains(response, 'id="requestor-search"')
        self.assertContains(response, 'id="assignee-search"')
        self.assertContains(response, 'list="assignee-options"')

    def test_user_sees_only_related_tasks(self):
        related = self.make_task()
        self.make_task(title="其他部門工作", request_department=self.other_department, requestor=self.outsider, owner_department=self.other_department, primary_assignee=self.outsider)
        self.assertEqual(list(visible_tasks(self.member, Task.objects.all())), [related])

    def test_manager_sees_department_scope(self):
        task = self.make_task()
        self.assertIn(task, visible_tasks(self.manager, Task.objects.all()))

    def test_system_settings_uses_custom_page_and_is_staff_only(self):
        self.client.force_login(self.manager)
        self.assertEqual(self.client.get(reverse("system_settings")).status_code, 403)
        admin = User.objects.create_superuser(
            "settings_admin", password="Strong-pass-123", display_name="設定管理員", role=User.Role.ADMIN
        )
        self.client.force_login(admin)
        response = self.client.get(reverse("system_settings"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "工作標籤與顏色")
        self.client.post(reverse("system_settings"), {
            "action": "save_tag", "name": "自訂標籤", "color": "#12AB34", "is_active": "on",
        })
        self.assertTrue(Tag.objects.filter(name="自訂標籤", color="#12AB34").exists())

    def test_calendar_shows_month_tasks_within_user_scope(self):
        today = timezone.localdate()
        visible = self.make_task(title="行事曆可見工作", due_date=today)
        self.make_task(
            title="行事曆其他部門工作",
            request_department=self.other_department,
            requestor=self.outsider,
            owner_department=self.other_department,
            primary_assignee=self.outsider,
            due_date=today,
            created_by=self.outsider,
            updated_by=self.outsider,
        )
        self.client.force_login(self.manager)
        response = self.client.get(reverse("task_calendar"), {"year": today.year, "month": today.month})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, visible.title)
        self.assertNotContains(response, "行事曆其他部門工作")
        self.assertContains(response, "今天")

    def test_excel_export_supports_selected_rows_fields_and_status_sheets(self):
        selected = self.make_task(title="選取匯出工作", progress=35)
        hidden = self.make_task(
            title="無權限匯出工作",
            request_department=self.other_department,
            requestor=self.outsider,
            owner_department=self.other_department,
            primary_assignee=self.outsider,
            created_by=self.outsider,
            updated_by=self.outsider,
        )
        self.client.force_login(self.manager)
        response = self.client.post(reverse("task_excel_export"), {
            "export_mode": "selected",
            "selected_ids": f"{selected.pk},{hidden.pk}",
            "split_by_status": "on",
            "fields": ["task_no", "title", "status", "progress", "due_date"],
        })
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        workbook = load_workbook(BytesIO(response.content))
        self.assertEqual(workbook.sheetnames, ["未開始"])
        sheet = workbook["未開始"]
        self.assertEqual([sheet.cell(4, column).value for column in range(1, 6)], ["工作編號", "工作名稱", "狀態", "進度", "預計完成日"])
        self.assertEqual(sheet.cell(5, 2).value, selected.title)
        self.assertEqual(sheet.cell(5, 4).value, 0.35)
        self.assertNotIn(hidden.title, [cell.value for row in sheet.iter_rows() for cell in row])

    def test_unauthorized_detail_returns_404(self):
        task = self.make_task()
        self.client.force_login(self.outsider)
        self.assertEqual(self.client.get(reverse("task_detail", args=[task.pk])).status_code, 404)

    def test_authorized_user_can_edit_directly_on_task_detail_page(self):
        task = self.make_task()
        self.client.force_login(self.manager)
        response = self.client.get(reverse("task_detail", args=[task.pk]))
        self.assertContains(response, 'id="inline-task-editor"')
        response = self.client.post(reverse("task_detail", args=[task.pk]), {
            "title": "詳細頁直接修改",
            "description": task.description,
            "project": "",
            "category": self.category.pk,
            "tags": [],
            "request_department": self.department.pk,
            "requestor": self.member.pk,
            "request_date": task.request_date,
            "owner_department": self.department.pk,
            "primary_assignee": self.member.pk,
            "collaborators": [],
            "priority": task.priority,
            "status": task.status,
            "progress": task.progress,
            "planned_start_date": "",
            "due_date": task.due_date,
            "waiting_reason": "",
            "cancellation_reason": "",
        })
        self.assertRedirects(response, reverse("task_detail", args=[task.pk]))
        task.refresh_from_db()
        self.assertEqual(task.title, "詳細頁直接修改")

    def test_viewer_cannot_create(self):
        viewer = User.objects.create_user("viewer", password="Strong-pass-123", display_name="唯讀", role=User.Role.VIEWER, department=self.department)
        self.client.force_login(viewer)
        self.assertEqual(self.client.get(reverse("task_create")).status_code, 403)

    def test_project_board_is_available(self):
        self.client.force_login(self.manager)
        self.assertEqual(self.client.get(reverse("project_board")).status_code, 200)

    def test_project_quick_entry_creates_multiple_tasks(self):
        project = Project.objects.create(
            name="快速輸入專案", owner_department=self.department, manager=self.manager,
            start_date=timezone.localdate(), due_date=timezone.localdate() + timedelta(days=30), created_by=self.manager,
        )
        self.client.force_login(self.manager)
        response = self.client.post(reverse("project_bulk_task_create", args=[project.pk]), {
            "title": ["第一項工作", "第二項工作", ""],
            "description": ["第一項說明", "第二項說明", ""],
            "assignee": [str(self.member.pk), str(self.member.pk), str(self.member.pk)],
            "due_date": [str(project.due_date), str(project.due_date), str(project.due_date)],
            "priority": [Task.Priority.HIGH, Task.Priority.MEDIUM, Task.Priority.LOW],
        })
        self.assertEqual(response.status_code, 302)
        self.assertEqual(Task.objects.filter(project=project).count(), 2)
        self.assertEqual(Task.objects.get(project=project, title="第一項工作").description, "第一項說明")

    def test_quick_status_advances_and_completed_task_moves_to_history(self):
        task = self.make_task()
        self.client.force_login(self.manager)
        first = self.client.post(reverse("task_quick_status", args=[task.pk]))
        self.assertEqual(first.json()["status"], Task.Status.IN_PROGRESS)
        second = self.client.post(reverse("task_quick_status", args=[task.pk]))
        self.assertTrue(second.json()["moved_to_history"])
        task.refresh_from_db()
        self.assertEqual(task.status, Task.Status.COMPLETED)
        self.assertNotContains(self.client.get(reverse("task_list")), task.title)
        self.assertContains(self.client.get(reverse("task_history")), task.title)

    def test_project_quick_status_returns_updated_project_summary(self):
        project = Project.objects.create(
            name="狀態同步專案", owner_department=self.department, manager=self.manager,
            start_date=timezone.localdate(), due_date=timezone.localdate() + timedelta(days=30),
            created_by=self.manager,
        )
        task = self.make_task(project=project)
        self.client.force_login(self.manager)
        first = self.client.post(reverse("task_quick_status", args=[task.pk])).json()
        self.assertEqual(first["project"]["progress"], 1)
        second = self.client.post(reverse("task_quick_status", args=[task.pk])).json()
        self.assertEqual(second["project"]["progress"], 100)
        self.assertEqual(second["project"]["completed_count"], 1)
        response = self.client.get(reverse("project_board"), {"project": project.pk})
        self.assertNotContains(response, task.title)

    def test_quick_status_can_select_waiting_and_return_to_in_progress(self):
        task = self.make_task(status=Task.Status.IN_PROGRESS, progress=30)
        self.client.force_login(self.manager)
        waiting = self.client.post(reverse("task_quick_status", args=[task.pk]), {
            "target_status": Task.Status.WAITING,
            "waiting_reason": "等待供應商回覆",
        })
        self.assertEqual(waiting.status_code, 200)
        self.assertEqual(waiting.json()["status"], Task.Status.WAITING)
        task.refresh_from_db()
        self.assertEqual(task.waiting_reason, "等待供應商回覆")
        resumed = self.client.post(reverse("task_quick_status", args=[task.pk]), {
            "target_status": Task.Status.IN_PROGRESS,
        })
        self.assertEqual(resumed.json()["status"], Task.Status.IN_PROGRESS)
        task.refresh_from_db()
        self.assertEqual(task.waiting_reason, "")

    def test_issue_note_moves_through_four_board_stages(self):
        self.client.force_login(self.member)
        response = self.client.post(reverse("issue_note_board"), {
            "reported_by_name": "現場同仁", "title": "系統欄位不清楚", "content": "需要確認操作方式",
        })
        self.assertEqual(response.status_code, 302)
        note = IssueNote.objects.get(title="系統欄位不清楚")
        self.assertEqual(note.status, IssueNote.Status.INBOX)
        self.client.post(reverse("issue_note_toggle", args=[note.pk]))
        self.client.post(reverse("issue_note_toggle", args=[note.pk]))
        self.client.post(reverse("issue_note_toggle", args=[note.pk]))
        note.refresh_from_db()
        self.assertEqual(note.status, IssueNote.Status.DONE)
        self.assertNotContains(self.client.get(reverse("issue_note_board")), note.title)
        history = self.client.get(reverse("task_history"))
        self.assertContains(history, note.title)
        self.assertContains(history, "恢復至收集箱")

    def test_manager_can_multi_filter_tasks_by_section(self):
        section_a = Department.objects.create(code="QA-A", name="品質一課", unit_type=Department.UnitType.SECTION, parent=self.department)
        section_b = Department.objects.create(code="QA-B", name="品質二課", unit_type=Department.UnitType.SECTION, parent=self.department)
        member_a = User.objects.create_user("member_a", password="Strong-pass-123", display_name="一課人員", department=section_a)
        member_b = User.objects.create_user("member_b", password="Strong-pass-123", display_name="二課人員", department=section_b)
        task_a = self.make_task(title="一課工作", request_department=section_a, requestor=member_a, owner_department=section_a, primary_assignee=member_a)
        self.make_task(title="二課工作", request_department=section_b, requestor=member_b, owner_department=section_b, primary_assignee=member_b)
        self.client.force_login(self.manager)
        response = self.client.get(reverse("task_list"), {"departments": [str(section_a.pk)]})
        self.assertContains(response, task_a.title)
        self.assertNotContains(response, "二課工作")

    def test_manager_can_filter_a_group_without_including_its_parent_section(self):
        section = Department.objects.create(
            code="QA-S", name="品質技術課",
            unit_type=Department.UnitType.SECTION, parent=self.department,
        )
        group = Department.objects.create(
            code="QA-S-G", name="供應商管理組",
            unit_type=Department.UnitType.GROUP, parent=section,
        )
        section_member = User.objects.create_user(
            "section_member", password="Strong-pass-123",
            display_name="課級人員", department=section,
        )
        group_member = User.objects.create_user(
            "group_member", password="Strong-pass-123",
            display_name="組級人員", department=group,
        )
        section_task = self.make_task(
            title="課級工作", request_department=section, requestor=section_member,
            owner_department=section, primary_assignee=section_member,
        )
        group_task = self.make_task(
            title="組級工作", request_department=group, requestor=group_member,
            owner_department=group, primary_assignee=group_member,
        )
        self.client.force_login(self.manager)
        response = self.client.get(reverse("task_list"), {"departments": [str(group.pk)]})
        self.assertContains(response, group_task.title)
        self.assertNotContains(response, section_task.title)

    def test_selecting_a_section_includes_tasks_from_its_groups(self):
        section = Department.objects.create(
            code="QA-P", name="品質系統課",
            unit_type=Department.UnitType.SECTION, parent=self.department,
        )
        group = Department.objects.create(
            code="QA-P-G", name="實驗室",
            unit_type=Department.UnitType.GROUP, parent=section,
        )
        group_member = User.objects.create_user(
            "lab_member", password="Strong-pass-123",
            display_name="實驗室人員", department=group,
        )
        group_task = self.make_task(
            title="實驗室工作", request_department=group, requestor=group_member,
            owner_department=group, primary_assignee=group_member,
        )
        self.client.force_login(self.manager)
        response = self.client.get(reverse("task_list"), {"departments": [str(section.pk)]})
        self.assertContains(response, group_task.title)

    def test_department_filter_includes_descendants_and_builds_status_summary(self):
        section = Department.objects.create(
            code="QA-D-S", name="部門摘要課",
            unit_type=Department.UnitType.SECTION, parent=self.department,
        )
        group = Department.objects.create(
            code="QA-D-S-G", name="部門摘要組",
            unit_type=Department.UnitType.GROUP, parent=section,
        )
        group_member = User.objects.create_user(
            "summary_member", password="Strong-pass-123",
            display_name="摘要人員", department=group,
        )
        descendant_task = self.make_task(
            title="部門下級工作", request_department=group, requestor=group_member,
            owner_department=group, primary_assignee=group_member,
            status=Task.Status.IN_PROGRESS,
            due_date=timezone.localdate() - timedelta(days=1),
        )
        self.client.force_login(self.manager)
        response = self.client.get(reverse("task_list"), {"departments": [str(self.department.pk)]})
        self.assertContains(response, descendant_task.title)
        summary = response.context["selected_unit_summaries"][0]
        self.assertEqual(summary["in_progress"], 1)
        self.assertEqual(summary["overdue"], 1)

    def test_organization_filter_is_saved_as_the_users_default_view(self):
        section = Department.objects.create(
            code="QA-SAVED", name="固定檢視課",
            unit_type=Department.UnitType.SECTION, parent=self.department,
        )
        section_member = User.objects.create_user(
            "saved_member", password="Strong-pass-123",
            display_name="固定檢視人員", department=section,
        )
        saved_task = self.make_task(
            title="固定檢視工作", request_department=section, requestor=section_member,
            owner_department=section, primary_assignee=section_member,
        )
        other_task = self.make_task(title="不在固定檢視的工作")
        self.client.force_login(self.manager)
        self.client.get(reverse("task_list"), {
            "organization_filter_submitted": "1",
            "departments": [str(section.pk)],
        })
        response = self.client.get(reverse("task_list"))
        self.assertContains(response, saved_task.title)
        self.assertNotContains(response, other_task.title)
        self.assertTrue(self.manager.preferred_task_departments.filter(pk=section.pk).exists())

    def test_user_can_clear_saved_organization_filter(self):
        self.manager.preferred_task_departments.set([self.department])
        self.client.force_login(self.manager)
        self.client.get(reverse("task_list"), {"organization_filter_submitted": "1"})
        self.assertFalse(self.manager.preferred_task_departments.exists())

    def test_focusing_a_card_keeps_all_saved_cards_but_filters_the_task_list(self):
        section_a = Department.objects.create(
            code="QA-CARD-A", name="卡片甲課",
            unit_type=Department.UnitType.SECTION, parent=self.department,
        )
        section_b = Department.objects.create(
            code="QA-CARD-B", name="卡片乙課",
            unit_type=Department.UnitType.SECTION, parent=self.department,
        )
        member_a = User.objects.create_user("card_a", password="Strong-pass-123", display_name="甲課人員", department=section_a)
        member_b = User.objects.create_user("card_b", password="Strong-pass-123", display_name="乙課人員", department=section_b)
        task_a = self.make_task(title="甲課卡片工作", request_department=section_a, requestor=member_a, owner_department=section_a, primary_assignee=member_a)
        task_b = self.make_task(title="乙課卡片工作", request_department=section_b, requestor=member_b, owner_department=section_b, primary_assignee=member_b)
        self.client.force_login(self.manager)
        response = self.client.get(reverse("task_list"), {
            "organization_filter_submitted": "1",
            "departments": [str(section_a.pk), str(section_b.pk)],
            "focus_departments": [str(section_a.pk)],
        })
        self.assertEqual(len(response.context["selected_unit_summaries"]), 2)
        self.assertContains(response, task_a.title)
        self.assertNotContains(response, task_b.title)
        self.assertEqual(self.manager.preferred_task_departments.count(), 2)

    def test_multiple_cards_can_be_focused_together(self):
        section_a = Department.objects.create(
            code="QA-MULTI-A", name="多選甲課",
            unit_type=Department.UnitType.SECTION, parent=self.department,
        )
        section_b = Department.objects.create(
            code="QA-MULTI-B", name="多選乙課",
            unit_type=Department.UnitType.SECTION, parent=self.department,
        )
        member_a = User.objects.create_user("multi_a", password="Strong-pass-123", display_name="多選甲", department=section_a)
        member_b = User.objects.create_user("multi_b", password="Strong-pass-123", display_name="多選乙", department=section_b)
        task_a = self.make_task(title="多選甲工作", request_department=section_a, requestor=member_a, owner_department=section_a, primary_assignee=member_a)
        task_b = self.make_task(title="多選乙工作", request_department=section_b, requestor=member_b, owner_department=section_b, primary_assignee=member_b)
        self.client.force_login(self.manager)
        response = self.client.get(reverse("task_list"), {
            "organization_filter_submitted": "1",
            "departments": [str(section_a.pk), str(section_b.pk)],
            "focus_departments": [str(section_a.pk), str(section_b.pk)],
        })
        self.assertContains(response, task_a.title)
        self.assertContains(response, task_b.title)
        self.assertEqual(response.context["focus_department_ids"], {section_a.pk, section_b.pk})
